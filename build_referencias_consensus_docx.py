from pathlib import Path
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor, Twips


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Referencias_Consensus_Gemelo_Digital.docx"
SKILL_ROOT = Path(r"C:\Users\alexs\.codex\plugins\cache\openai-primary-runtime\documents\26.826.12353\skills\documents")
sys.path.insert(0, str(SKILL_ROOT / "scripts"))
from table_geometry import apply_table_geometry  # noqa: E402


NAVY = "17324D"
BLUE = "2E74B5"
DEEP_BLUE = "1F4D78"
TEAL = "167D8D"
INK = "1F2937"
MUTED = "667085"
PALE_BLUE = "E8EEF5"
PALE_TEAL = "EAF6F7"
PALE_GRAY = "F4F6F9"
WHITE = "FFFFFF"
GOLD = "B9852A"


REFERENCES = [
    {
        "number": "01",
        "short": "Gemelo digital de ciudad para gestión de desastres",
        "apa": "Fan, C., Zhang, C., Yahja, A., & Mostafavi, A. (2021). Disaster City Digital Twin: A vision for integrating artificial and human intelligence for disaster management. International Journal of Information Management, 56, 102049.",
        "doi": "https://doi.org/10.1016/j.ijinfomgt.2019.102049",
        "translated_title": "Gemelo digital de una ciudad en desastre: una visión para integrar inteligencia artificial y humana en la gestión de desastres.",
        "quote": "This paper presents the Disaster City Digital Twin as a unifying paradigm.",
        "translation": "Este artículo presenta el Gemelo Digital de Ciudad en Desastre como un paradigma unificador.",
        "paraphrase": "El gemelo digital urbano puede funcionar como una arquitectura integradora que articula datos situacionales, inteligencia artificial, decisiones de múltiples actores y análisis dinámico de redes durante una emergencia (Fan et al., 2021).",
        "use": "Marco conceptual y justificación de la arquitectura integral.",
        "consensus": "https://consensus.app/papers/disaster-city-digital-twin-a-vision-for-integrating-fan-zhang/a604ce54a0225e10857f6d96923536e3/",
    },
    {
        "number": "02",
        "short": "Bucles de información en sistemas urbanos inteligentes",
        "apa": "Ford, D. N., & Wolf, C. M. (2020). Smart cities with digital twin systems for disaster management. Journal of Management in Engineering, 36(4), 04020027.",
        "doi": "https://doi.org/10.1061/(ASCE)ME.1943-5479.0000779",
        "translated_title": "Ciudades inteligentes con sistemas de gemelo digital para la gestión de desastres.",
        "quote": "Information loops, as opposed to individual components, are identified as a critical future focus of SCDT development.",
        "translation": "Los bucles de información, más que los componentes aislados, se identifican como un foco futuro crítico del desarrollo de estos sistemas.",
        "paraphrase": "La eficacia de un gemelo digital urbano no depende solo de acumular sensores o modelos, sino de cerrar ciclos de información que conecten observación, simulación, decisión y actualización del sistema (Ford & Wolf, 2020).",
        "use": "Diseño del flujo bidireccional entre los niveles micro, meso y macro.",
        "consensus": "https://consensus.app/papers/smart-cities-with-digital-twin-systems-for-disaster-ford-wolf/cd330389946e5079bc0608e2a07469ae/",
    },
    {
        "number": "03",
        "short": "Vacíos en los gemelos digitales posdesastre",
        "apa": "Lagap, U., & Ghaffarian, S. (2024). Digital post-disaster risk management twinning: A review and improved conceptual framework. International Journal of Disaster Risk Reduction, 110, 104629.",
        "doi": "https://doi.org/10.1016/j.ijdrr.2024.104629",
        "translated_title": "Gemelización digital de la gestión del riesgo posdesastre: revisión y marco conceptual mejorado.",
        "quote": "There is a need for diverse data collection methods, considering demographic and financial aspects, understanding social dynamics, employing dynamic models, recognizing interconnected systems.",
        "translation": "Se requieren métodos diversos de recolección de datos, considerar aspectos demográficos y financieros, comprender las dinámicas sociales, usar modelos dinámicos y reconocer sistemas interconectados.",
        "paraphrase": "La literatura reciente señala que los gemelos posdesastre todavía deben integrar mejor las condiciones demográficas y económicas, las dinámicas sociales, la evolución temporal y las interdependencias entre sistemas (Lagap & Ghaffarian, 2024).",
        "use": "Formulación del vacío de conocimiento y variables socioeconómicas.",
        "consensus": "https://consensus.app/papers/digital-postdisaster-risk-management-twinning-a-review-lagap-ghaffarian/eca51f1ad3d15bc390f4751d70e46817/",
    },
    {
        "number": "04",
        "short": "Evacuación por inundación en Kibera con datos abiertos",
        "apa": "Butters, O., & Dawson, R. J. (2025). Flood evacuation in informal settlements: Application of an agent-based model to Kibera using open data. Urban Science, 9(1), 12.",
        "doi": "https://doi.org/10.3390/urbansci9010012",
        "translated_title": "Evacuación por inundación en asentamientos informales: aplicación de un modelo basado en agentes a Kibera mediante datos abiertos.",
        "quote": "This paper reports on the first application of an agent-based model to an informal settlement, Kibera.",
        "translation": "Este artículo informa la primera aplicación de un modelo basado en agentes a un asentamiento informal: Kibera.",
        "paraphrase": "El caso de Kibera demuestra que un modelo de evacuación basado en agentes puede adaptarse a un asentamiento informal con escasez de datos y utilizar fuentes abiertas para identificar refugios funcionales y rutas propensas a congestión (Butters & Dawson, 2025).",
        "use": "Antecedente central, viabilidad con datos abiertos y validación del escenario de inundación.",
        "consensus": "https://consensus.app/papers/flood-evacuation-in-informal-settlements-application-of-butters-dawson/68ce2d316108506dacdc2363b69f664e/",
    },
    {
        "number": "05",
        "short": "Recuperación de hogares formales e informales en Tacloban",
        "apa": "Ghaffarian, S., Roy, D., Filatova, T., & Kerle, N. (2021). Agent-based modelling of post-disaster recovery with remote sensing data. International Journal of Disaster Risk Reduction, 60, 102285.",
        "doi": "https://doi.org/10.1016/j.ijdrr.2021.102285",
        "translated_title": "Modelado basado en agentes de la recuperación posdesastre con datos de teledetección.",
        "quote": "Formal and informal (slum) sector households are differentiated in the model to explore their resilience and different recovery patterns.",
        "translation": "El modelo diferencia hogares de sectores formales e informales para explorar su resiliencia y sus distintos patrones de recuperación.",
        "paraphrase": "La simulación de la recuperación debe distinguir hogares formales e informales, porque sus recursos, formas de reconstrucción, acceso al empleo y decisiones de reubicación generan trayectorias de resiliencia diferentes (Ghaffarian et al., 2021).",
        "use": "Nivel micro, recuperación económica y calibración con teledetección.",
        "consensus": "https://consensus.app/papers/agentbased-modelling-of-postdisaster-recovery-with-ghaffarian-roy/7c654ec820cf5ad293f6d18dd370810a/",
    },
    {
        "number": "06",
        "short": "Gemelos urbanos centrados en las personas",
        "apa": "Ye, X., Du, J., Han, Y., Newman, G., Retchless, D., Zou, L., Ham, Y., & Cai, Z. (2023). Developing human-centered urban digital twins for community infrastructure resilience: A research agenda. Journal of Planning Literature, 38(2), 187-199.",
        "doi": "https://doi.org/10.1177/08854122221137861",
        "translated_title": "Desarrollo de gemelos digitales urbanos centrados en las personas para la resiliencia de infraestructura comunitaria: una agenda de investigación.",
        "quote": "This article highlights the need for integrating multi-agent interactions, artificial intelligence, and coupled natural-physical-social systems into a human-centered UDTs framework.",
        "translation": "El artículo destaca la necesidad de integrar interacciones multiagente, inteligencia artificial y sistemas naturales, físicos y sociales acoplados en un marco centrado en las personas.",
        "paraphrase": "Un gemelo digital orientado a la resiliencia comunitaria debe representar a las personas como participantes del sistema e integrar sus interacciones con la infraestructura, el ambiente y las instituciones, en lugar de limitarse a reproducir objetos físicos (Ye et al., 2023).",
        "use": "Justificación del enfoque humano, redes de apoyo e interacción entre escalas.",
        "consensus": "https://consensus.app/papers/developing-humancentered-urban-digital-twins-for-ye-du/bfbc3037ff005c1480c50ce0e9ac1e2a/",
    },
    {
        "number": "07",
        "short": "Gemelos urbanos y Marco de Sendai",
        "apa": "Macatulad, E., & Biljecki, F. (2024). Continuing from the Sendai Framework midterm: Opportunities for urban digital twins in disaster risk management. International Journal of Disaster Risk Reduction, 102, 104310.",
        "doi": "https://doi.org/10.1016/j.ijdrr.2024.104310",
        "translated_title": "Continuación desde la evaluación intermedia del Marco de Sendai: oportunidades de los gemelos digitales urbanos en la gestión del riesgo de desastres.",
        "quote": "Research directions show opportunities for urban digital twins in disaster risk management-particularly, as the integrating framework and platform.",
        "translation": "Las líneas de investigación muestran oportunidades para los gemelos digitales urbanos, particularmente como marco y plataforma de integración.",
        "paraphrase": "Debido a que el riesgo urbano es multidimensional, multiescala, multiactor y multiamenaza, los gemelos digitales ofrecen una plataforma pertinente para integrar procesos de gestión del riesgo coherentes con las prioridades del Marco de Sendai (Macatulad & Biljecki, 2024).",
        "use": "Marco normativo, enfoque multiamenaza y defensa de la arquitectura multiescala.",
        "consensus": "https://consensus.app/papers/continuing-from-the-sendai-framework-midterm-macatulad-biljecki/e544a243e95952b0ab99c69cbbe1caf1/",
    },
    {
        "number": "08",
        "short": "Evaluación de resiliencia comunitaria mediante agentes",
        "apa": "Aghababaei, M., & Koliou, M. (2023). Community resilience assessment via agent-based modeling approach. Computer-Aided Civil and Infrastructure Engineering, 38(7), 920-939.",
        "doi": "https://doi.org/10.1111/mice.12916",
        "translated_title": "Evaluación de la resiliencia comunitaria mediante un enfoque de modelado basado en agentes.",
        "quote": "The results show the capability of the proposed computational model to capture the complex behaviors and dynamics of interdependent systems within a community.",
        "translation": "Los resultados muestran la capacidad del modelo computacional para capturar comportamientos complejos y dinámicas de sistemas interdependientes dentro de una comunidad.",
        "paraphrase": "Los modelos basados en agentes permiten representar hogares, personas, servicios de salud, educación, empresas y redes de utilidad como sistemas interdependientes, y comparar estrategias de mitigación antes de un desastre (Aghababaei & Koliou, 2023).",
        "use": "Diseño de agentes, métricas de resiliencia y evaluación de intervenciones.",
        "consensus": "https://consensus.app/papers/community-resilience-assessment-via-agent%E2%80%90based-aghababaei-koliou/2eb409631c1058de975ae9026333280c/",
    },
    {
        "number": "09",
        "short": "Infraestructuras críticas como sistemas sociotécnicos",
        "apa": "Gebhard, T., Sattler, B. J., Gunkel, J., Marquard, M., & Tundis, A. (2025). Improving the resilience of socio-technical urban critical infrastructures with digital twins: Challenges, concepts, and modeling. Sustainability Analytics and Modeling, 5, 100036.",
        "doi": "https://doi.org/10.1016/j.samod.2024.100036",
        "translated_title": "Mejora de la resiliencia de infraestructuras críticas urbanas sociotécnicas con gemelos digitales: desafíos, conceptos y modelado.",
        "quote": "Our approach frames CIs as socio-technical systems and integrates the human perspective into the modeling process and resilience assessment.",
        "translation": "Nuestro enfoque concibe las infraestructuras críticas como sistemas sociotécnicos e integra la perspectiva humana en el modelado y la evaluación de resiliencia.",
        "paraphrase": "La evaluación integral de resiliencia exige modelar conjuntamente las redes de energía, agua y transporte, sus dependencias y la forma en que la población utiliza o modifica su demanda durante la crisis (Gebhard et al., 2025).",
        "use": "Nivel macro, capacidad institucional e interdependencia de infraestructura crítica.",
        "consensus": "https://consensus.app/papers/improving-the-resilience-of-sociotechnical-urban-gebhard-sattler/ea8b302f7a755b86bf5e4f93837179e0/",
    },
    {
        "number": "10",
        "short": "Colaboración y participación en gemelos digitales urbanos",
        "apa": "Therias, A., & Rafiee, A. (2023). City digital twins for urban resilience. International Journal of Digital Earth, 16(2), 4164-4190.",
        "doi": "https://doi.org/10.1080/17538947.2023.2264827",
        "translated_title": "Gemelos digitales de ciudad para la resiliencia urbana.",
        "quote": "The paper concludes by emphasizing the important role of stakeholders in shaping CDTs that can be successfully integrated by the communities they serve.",
        "translation": "El artículo concluye destacando el papel de los actores interesados en configurar gemelos digitales que las comunidades puedan integrar con éxito.",
        "paraphrase": "La utilidad social de un gemelo digital depende de incorporar a comunidades, autoridades y otros actores en su diseño, de modo que la plataforma favorezca comprensión compartida, exploración de escenarios y coordinación colectiva (Therias & Rafiee, 2023).",
        "use": "Validación participativa, gobernanza y adopción comunitaria.",
        "consensus": "https://consensus.app/papers/city-digital-twins-for-urban-resilience-therias-rafiee/c85e36ae7c9652c0a8f0eb5015589e1a/",
    },
]


OBJECTIVES = [
    ("Arquitectura anidada micro-meso-macro", "Fan et al. (2021); Ye et al. (2023); Macatulad y Biljecki (2024)"),
    ("Integración de datos físicos, sociales e institucionales", "Lagap y Ghaffarian (2024); Gebhard et al. (2025)"),
    ("Evacuación y recuperación mediante simulación", "Butters y Dawson (2025); Ghaffarian et al. (2021)"),
    ("Sistemas comunitarios e infraestructuras interdependientes", "Aghababaei y Koliou (2023); Ford y Wolf (2020)"),
    ("Validación participativa y adopción social", "Therias y Rafiee (2023); Ye et al. (2023)"),
]


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill, left_border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if left_border:
        borders = p_pr.find(qn("w:pBdr"))
        if borders is None:
            borders = OxmlElement("w:pBdr")
            p_pr.append(borders)
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), left_border)
        borders.append(left)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, color="D9E2EC", size="6"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{side}"))
        if tag is None:
            tag = OxmlElement(f"w:{side}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_hyperlink(paragraph, text, url, color=BLUE, underline=True):
    part = paragraph.part
    rel_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    color_el = OxmlElement("w:color")
    color_el.set(qn("w:val"), color)
    r_pr.append(color_el)
    if underline:
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        r_pr.append(u)
    size = OxmlElement("w:sz")
    size.set(qn("w:val"), "20")
    r_pr.append(size)
    new_run.append(r_pr)
    text_el = OxmlElement("w:t")
    text_el.text = text
    new_run.append(text_el)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_field(paragraph, instruction):
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    display = OxmlElement("w:t")
    display.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, display, fld_end])
    set_run_font(run, size=9, color=MUTED)


def set_keep_with_next(paragraph, value=True):
    paragraph.paragraph_format.keep_with_next = value


def add_labeled_paragraph(doc, label, text, *, italic=False, after=6, hanging=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if hanging:
        p.paragraph_format.left_indent = Inches(0.35)
        p.paragraph_format.first_line_indent = Inches(-0.35)
    label_run = p.add_run(label)
    set_run_font(label_run, size=10.5, color=DEEP_BLUE, bold=True)
    body = p.add_run(text)
    set_run_font(body, size=10.5, color=INK, italic=italic)
    return p


def add_quote_block(doc, quote, translation):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.right_indent = Inches(0.12)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.18
    shade_paragraph(p, PALE_TEAL, TEAL)
    label = p.add_run("FRAGMENTO TEXTUAL · RESUMEN\n")
    set_run_font(label, size=8.5, color=TEAL, bold=True)
    original = p.add_run(f'“{quote}”')
    set_run_font(original, size=10.5, color=NAVY, bold=True, italic=True)
    tr_label = p.add_run("\nTraducción propia: ")
    set_run_font(tr_label, size=9.5, color=TEAL, bold=True)
    tr = p.add_run(f'“{translation}”')
    set_run_font(tr, size=9.5, color=INK, italic=True)
    return p


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DEEP_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True


def configure_section(section):
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    hp = header.paragraphs[0]
    hp.paragraph_format.space_after = Pt(0)
    hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5))
    left = hp.add_run("GEMELO DIGITAL DE RESILIENCIA")
    set_run_font(left, size=8.5, color=MUTED, bold=True)
    right = hp.add_run("\tMATRIZ DE EVIDENCIA · CONSENSUS")
    set_run_font(right, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    lead = fp.add_run("Página ")
    set_run_font(lead, size=9, color=MUTED)
    add_field(fp, "PAGE")
    mid = fp.add_run(" de ")
    set_run_font(mid, size=9, color=MUTED)
    add_field(fp, "NUMPAGES")


def add_cover(doc):
    for _ in range(3):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(10)

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(18)
    r = kicker.add_run("MATRIZ DE EVIDENCIA CIENTÍFICA")
    set_run_font(r, size=10, color=TEAL, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(12)
    title.paragraph_format.keep_with_next = True
    r = title.add_run("10 referencias académicas\nlocalizadas mediante Consensus")
    set_run_font(r, size=29, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(28)
    r = subtitle.add_run("Fragmentos traducidos y paráfrasis para la escritura del artículo científico")
    set_run_font(r, size=13, color=DEEP_BLUE, italic=True)

    project = doc.add_paragraph()
    project.alignment = WD_ALIGN_PARAGRAPH.CENTER
    project.paragraph_format.left_indent = Inches(0.35)
    project.paragraph_format.right_indent = Inches(0.35)
    project.paragraph_format.space_after = Pt(26)
    shade_paragraph(project, PALE_BLUE)
    r = project.add_run("PROYECTO\n")
    set_run_font(r, size=8.5, color=BLUE, bold=True)
    r = project.add_run("Gemelo Digital Poblacional para Evaluación de Resiliencia ante Desastres Naturales en Asentamientos Informales: Arquitectura Multi-escala para Ciudades del Sur Global")
    set_run_font(r, size=11.5, color=NAVY, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(8)
    r = meta.add_run("31 de agosto de 2026  ·  10 artículos revisados  ·  DOI verificados")
    set_run_font(r, size=10, color=MUTED, bold=True)

    search = doc.add_paragraph()
    search.alignment = WD_ALIGN_PARAGRAPH.CENTER
    search.paragraph_format.left_indent = Inches(0.35)
    search.paragraph_format.right_indent = Inches(0.35)
    search.paragraph_format.space_after = Pt(4)
    r = search.add_run("Consulta: “Digital twins and agent-based models for disaster resilience, evacuation, social vulnerability, informal settlements, and Global South cities”.")
    set_run_font(r, size=9.5, color=MUTED, italic=True)

    source = doc.add_paragraph()
    source.alignment = WD_ALIGN_PARAGRAPH.CENTER
    source.paragraph_format.space_after = Pt(0)
    add_hyperlink(source, "Abrir búsqueda en Consensus", "https://consensus.app/search/digital-twins-agent-models-disaster-resilience/wihb1m19Qwub9vYjM_aHZA/", color=TEAL)
    doc.add_page_break()


def add_method(doc):
    doc.add_heading("Cómo utilizar esta matriz", level=1)
    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    intro.paragraph_format.line_spacing = 1.25
    intro.add_run("Cada ficha separa cuatro niveles de uso académico: ")
    bold = intro.add_run("referencia APA 7, fragmento original, traducción propia y paráfrasis.")
    bold.bold = True
    intro.add_run(" La paráfrasis puede integrarse al manuscrito con su cita, pero no debe presentarse como un resultado producido por el proyecto.")

    note = doc.add_paragraph()
    note.paragraph_format.left_indent = Inches(0.22)
    note.paragraph_format.right_indent = Inches(0.12)
    note.paragraph_format.space_after = Pt(14)
    shade_paragraph(note, PALE_GRAY, GOLD)
    label = note.add_run("NOTA DE INTEGRIDAD ACADÉMICA\n")
    set_run_font(label, size=8.5, color=GOLD, bold=True)
    body = note.add_run("Los fragmentos se tomaron del resumen de cada publicación y se mantienen breves. Toda traducción se identifica como propia y toda paráfrasis conserva su cita.")
    set_run_font(body, size=10.5, color=INK)

    doc.add_heading("Mapa de lectura", level=2)
    for label, text in (
        ("Escala micro", "hogares, evacuación, recuperación y vulnerabilidad diferencial."),
        ("Escala meso", "redes comunitarias, participación, servicios y coordinación local."),
        ("Escala macro", "infraestructura crítica, capacidad institucional e integración urbana."),
    ):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.left_indent = Inches(0.375)
        p.paragraph_format.first_line_indent = Inches(-0.188)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=10.5, color=DEEP_BLUE, bold=True)
        r = p.add_run(text)
        set_run_font(r, size=10.5, color=INK)
    doc.add_page_break()


def add_reference_entry(doc, item):
    heading = doc.add_heading(f"{item['number']}  {item['short']}", level=1)
    heading.paragraph_format.page_break_before = False
    set_keep_with_next(heading)

    p = add_labeled_paragraph(doc, "Referencia APA 7.  ", item["apa"], after=3, hanging=True)
    add_hyperlink(p, "  DOI", item["doi"], color=TEAL)

    add_labeled_paragraph(doc, "Título traducido.  ", item["translated_title"], italic=True, after=7)
    add_quote_block(doc, item["quote"], item["translation"])

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(7)
    p.paragraph_format.line_spacing = 1.25
    label = p.add_run("PARÁFRASIS LISTA PARA EL ARTÍCULO\n")
    set_run_font(label, size=8.5, color=BLUE, bold=True)
    text = p.add_run(item["paraphrase"])
    set_run_font(text, size=10.5, color=INK)

    use = doc.add_paragraph()
    use.paragraph_format.space_after = Pt(3)
    use.paragraph_format.keep_together = True
    lab = use.add_run("Uso recomendado  ")
    set_run_font(lab, size=9.5, color=TEAL, bold=True)
    val = use.add_run(item["use"])
    set_run_font(val, size=9.5, color=INK)

    links = doc.add_paragraph()
    links.paragraph_format.space_after = Pt(12)
    lab = links.add_run("Fuentes  ")
    set_run_font(lab, size=9.5, color=MUTED, bold=True)
    add_hyperlink(links, "Registro en Consensus", item["consensus"], color=TEAL)
    sep = links.add_run("  ·  ")
    set_run_font(sep, size=9.5, color=MUTED)
    add_hyperlink(links, "Publicación / DOI", item["doi"], color=TEAL)


def add_synthesis(doc):
    doc.add_page_break()
    doc.add_heading("Síntesis integradora sugerida", level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(
        "La literatura revisada permite sostener que un gemelo digital para la resiliencia de asentamientos informales no debe limitarse a una representación tridimensional del entorno construido. Su arquitectura debe cerrar bucles entre observación, simulación y decisión; combinar comportamiento de hogares, redes comunitarias e infraestructura crítica; y reconocer que los patrones de evacuación y recuperación varían según condiciones socioeconómicas y territoriales (Ford & Wolf, 2020; Ghaffarian et al., 2021; Ye et al., 2023)."
    )
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.25
    p.add_run(
        "Los estudios de Kibera y Tacloban demuestran además que los modelos basados en agentes y la teledetección pueden producir conocimiento operativo en contextos del Sur Global con información incompleta (Butters & Dawson, 2025; Ghaffarian et al., 2021). En conjunto, estas evidencias respaldan una arquitectura humana, multiescala, multiactor y multiamenaza, alineada con el Marco de Sendai y validada de manera participativa (Macatulad & Biljecki, 2024; Therias & Rafiee, 2023)."
    )

    doc.add_heading("Correspondencia con los objetivos del proyecto", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    header = table.rows[0].cells
    header[0].text = "Objetivo del proyecto"
    header[1].text = "Referencias de respaldo"
    set_repeat_table_header(table.rows[0])
    for cell in header:
        set_cell_shading(cell, PALE_BLUE)
        set_cell_border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(0)
            for r in p.runs:
                set_run_font(r, size=9.5, color=NAVY, bold=True)

    for objective, refs in OBJECTIVES:
        cells = table.add_row().cells
        cells[0].text = objective
        cells[1].text = refs
        for idx, cell in enumerate(cells):
            set_cell_border(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.15
                for r in p.runs:
                    set_run_font(r, size=9.5, color=INK, bold=(idx == 0))
    apply_table_geometry(table, [5000, 4360], table_width_dxa=9360, indent_dxa=120, cell_margins_dxa={"top": 100, "bottom": 100, "start": 120, "end": 120})

    source = doc.add_paragraph()
    source.paragraph_format.space_before = Pt(4)
    source.paragraph_format.space_after = Pt(12)
    r = source.add_run("Fuente: elaboración propia a partir de los diez artículos localizados mediante Consensus.")
    set_run_font(r, size=8.5, color=MUTED, italic=True)

    doc.add_heading("Recomendación de integridad académica", level=1)
    note = doc.add_paragraph()
    note.paragraph_format.left_indent = Inches(0.22)
    note.paragraph_format.right_indent = Inches(0.12)
    note.paragraph_format.space_after = Pt(0)
    shade_paragraph(note, PALE_GRAY, GOLD)
    r = note.add_run(
        "Antes del envío del manuscrito, compruébese la paginación final y el estilo específico de la revista. Las citas textuales deben conservar comillas y localización; las paráfrasis deben mantener su cita. Una traducción textual debe identificarse siempre como traducción propia."
    )
    set_run_font(r, size=10.5, color=INK)


def add_document_metadata(doc):
    props = doc.core_properties
    props.title = "10 referencias académicas localizadas mediante Consensus"
    props.subject = "Gemelo digital poblacional y resiliencia ante desastres en asentamientos informales"
    props.keywords = "Consensus, gemelo digital, resiliencia, desastres, asentamientos informales, modelos basados en agentes"
    props.comments = "Matriz académica con fragmentos breves, traducciones propias y paráfrasis en español."
    props.author = "Equipo del proyecto Gemelo Digital de Resiliencia"


def build():
    doc = Document()
    configure_styles(doc)
    for section in doc.sections:
        configure_section(section)
    add_document_metadata(doc)
    add_cover(doc)
    add_method(doc)
    for index, item in enumerate(REFERENCES):
        add_reference_entry(doc, item)
        if index in {1, 3, 5, 7}:
            doc.add_page_break()
    add_synthesis(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
